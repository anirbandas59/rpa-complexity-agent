"""DOCX document parser using python-docx.

Extracts text with heading-based structure, tables, and metadata from DOCX documents.
Provides schema-compatible output with parse_pdf for downstream consistency.
"""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

from config.logging_config import get_logger
from core.exceptions import DocumentProcessingError
from core.models.document import ParsedDocument

logger = get_logger("docx_parser")

# Standard estimate: 250 words per page
WORDS_PER_PAGE_ESTIMATE = 250


def parse_docx(file_path: str) -> ParsedDocument:
    """
    Parse a DOCX document and extract text, tables, and metadata.

    STEP 1: File validation
    STEP 2: Open document with python-docx
    STEP 3: Extract metadata from core_properties
    STEP 4: Extract text with heading-based structure markers
    STEP 5: Extract tables
    STEP 6: Estimate page count
    STEP 7: Build and return ParsedDocument

    Args:
        file_path: Path to the DOCX file to parse

    Returns:
        ParsedDocument: Parsed document model with all extracted content

    Raises:
        DocumentProcessingError: If file validation fails or opening fails
    """
    path = Path(file_path)

    # STEP 1: FILE VALIDATION
    _validate_docx_file(path)

    doc = None
    try:
        # STEP 2: OPEN DOCUMENT
        try:
            doc = Document(str(path))
            logger.info(f"Opened DOCX: {path}")
        except Exception as e:
            raise DocumentProcessingError(
                message=f"Failed to open DOCX: {path}",
                context={"error": str(e), "file_path": str(path)},
            )

        # STEP 3: EXTRACT METADATA
        metadata = _extract_metadata(doc, path)

        # STEP 4: EXTRACT TEXT WITH STRUCTURE
        full_text, extraction_warnings = _extract_text(doc)

        # STEP 5: EXTRACT TABLES
        tables = _extract_tables(doc, extraction_warnings)

        # STEP 6: ESTIMATE PAGE COUNT
        word_count = len(full_text.split())
        estimated_pages = max(1, word_count // WORDS_PER_PAGE_ESTIMATE)
        metadata["page_count_estimated"] = "true"

        # STEP 7: BUILD AND RETURN ParsedDocument
        parsed_doc = ParsedDocument(
            source_path=str(path.absolute()),
            file_type="docx",
            full_text=full_text,
            tables=tables,
            page_count=estimated_pages,
            metadata=metadata,
            extraction_warnings=extraction_warnings,
        )

        logger.info(
            f"Successfully parsed DOCX: {path} "
            f"(~{estimated_pages} pages, {word_count} words)"
        )
        return parsed_doc

    finally:
        # No explicit close needed for python-docx Documents
        pass


def _validate_docx_file(path: Path) -> None:
    """
    Validate DOCX file before processing.

    Checks for:
    - File existence
    - DOCX extension (.docx, case-insensitive)
    - Non-empty file

    Args:
        path: Path object to validate

    Raises:
        DocumentProcessingError: If any validation check fails
    """
    # Check existence
    if not path.exists():
        raise DocumentProcessingError(
            message=f"DOCX file not found: {path}",
            context={"file_path": str(path)},
        )

    # Check extension
    if path.suffix.lower() != ".docx":
        raise DocumentProcessingError(
            message=f"File is not a DOCX: {path}",
            context={"file_path": str(path), "extension": path.suffix},
        )

    # Check file size
    file_size = os.path.getsize(str(path))
    if file_size == 0:
        raise DocumentProcessingError(
            message=f"DOCX file is empty: {path}",
            context={"file_path": str(path)},
        )

    logger.debug(f"File validation passed: {path} ({file_size} bytes)")


def _extract_metadata(doc: Any, path: Path) -> dict[str, str]:
    """
    Extract metadata from DOCX document core properties.

    Extracts author, title, subject, created, modified, last_modified_by.
    Converts datetime objects to ISO format strings.
    Adds file-level metadata: file_size_bytes, paragraph_count, table_count.
    All values are converted to strings; None values are removed.

    Args:
        doc: python-docx Document object
        path: Path to the DOCX file

    Returns:
        Dictionary with metadata (all string values)
    """
    metadata = {}

    # Extract core properties
    core_props = doc.core_properties
    for attr in [
        "author",
        "title",
        "subject",
        "created",
        "modified",
        "last_modified_by",
    ]:
        value = getattr(core_props, attr, None)
        cleaned = _format_metadata_value(value)
        if cleaned is not None:
            metadata[attr] = cleaned

    # Add file-level metadata
    metadata["file_size_bytes"] = str(os.path.getsize(str(path)))
    metadata["paragraph_count"] = str(len(doc.paragraphs))
    metadata["table_count"] = str(len(doc.tables))

    logger.debug(f"Extracted metadata: {len(metadata)} fields")
    return metadata


def _format_metadata_value(value: Any) -> str | None:
    """
    Format a core_properties value to string.

    Converts datetime to ISO format.
    Returns None for None values.
    Converts other types to string.

    Args:
        value: Raw metadata value (may be None, datetime, or string)

    Returns:
        Formatted string or None if value is None
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    stripped = str(value).strip()
    return stripped if stripped else None


def _extract_text(doc: Any) -> tuple[str, list[str]]:
    """
    Extract text from DOCX document with heading-based structure markers.

    Uses heading levels to insert structural markers instead of page numbers.
    Heading 1 → "=== text ==="
    Heading 2 → "--- text ---"
    Heading 3 → "  > text"
    Other → plain text with newline

    Args:
        doc: python-docx Document object

    Returns:
        Tuple of (full_text, extraction_warnings)
    """
    text_parts = []
    warnings = []
    last_was_heading = False

    for para in doc.paragraphs:
        text = para.text.strip()

        # Classify heading level
        level = _get_paragraph_level(para.style.name)

        if level > 0:
            # This is a heading
            if level == 1:
                marker = f"\n=== {text} ===\n"
            elif level == 2:
                marker = f"\n--- {text} ---\n"
            else:  # level == 3
                marker = f"\n  > {text}\n"
            text_parts.append(marker)
            last_was_heading = True
        elif text:
            # This is body text (non-empty)
            text_parts.append(text + "\n")
            last_was_heading = False
        elif last_was_heading:
            # Preserve one blank line after headings
            text_parts.append("\n")
            last_was_heading = False

    full_text = "".join(text_parts)

    # Track warning if document is too short
    if len(full_text) < 100:
        warnings.append(
            "Document contains very little text — "
            "may be image-based or template-only"
        )

    logger.debug(
        f"Extracted text: {len(full_text)} characters, {len(warnings)} warnings"
    )
    return full_text, warnings


def _get_paragraph_level(style_name: str) -> int:
    """
    Get the heading level from a paragraph style name.

    Case-insensitive matching for "Heading 1", "Heading 2", "Heading 3".

    Args:
        style_name: The style name from paragraph.style.name

    Returns:
        Heading level (1, 2, or 3) or 0 for non-heading styles
    """
    if not style_name:
        return 0

    lower_name = style_name.lower()

    if "heading 1" in lower_name:
        return 1
    if "heading 2" in lower_name:
        return 2
    if "heading 3" in lower_name:
        return 3

    return 0


def _extract_tables(doc: Any, warnings: list[str]) -> list[list[dict[str, str]]]:
    """
    Extract tables from DOCX document.

    For each table, converts rows to list of dicts using first row as headers.
    Handles merged cells via python-docx cell.text API.
    Skips empty rows and tables with no data rows.
    Failures add warnings and continue.

    Args:
        doc: python-docx Document object
        warnings: List to append warnings to (modified in place)

    Returns:
        List of tables, each table is a list of row dicts
    """
    tables = []

    for table_idx, table in enumerate(doc.tables):
        try:
            table_rows = _extract_table(table)
            if table_rows:  # Only add non-empty tables
                tables.append(table_rows)
        except Exception as e:
            warnings.append(f"Table {table_idx + 1} extraction failed: {e}")
            logger.warning(f"Table {table_idx + 1} extraction failed: {e}")

    logger.debug(f"Extracted {len(tables)} tables")
    return tables


def _extract_table(table: Any) -> list[dict[str, str]]:
    """
    Extract a single python-docx table to list of dicts.

    First row becomes column headers. Remaining rows become dicts.
    Replaces empty cells with empty string.
    Returns empty list if table has <= 1 row (headers only).

    Args:
        table: python-docx Table object

    Returns:
        List of dicts (rows), empty list if no data rows
    """
    if not table.rows:
        return []

    # Extract headers from first row
    headers = []
    for cell in table.rows[0].cells:
        header = cell.text.strip()
        headers.append(header if header else "")

    # If only one row (headers only), return empty list
    if len(table.rows) <= 1:
        return []

    # Convert remaining rows to dicts
    result = []
    for row in table.rows[1:]:
        # Skip rows where all cells are empty
        row_text = [cell.text.strip() for cell in row.cells]
        if not any(row_text):
            continue

        # Build dict for this row
        row_dict = {}
        for col_idx, header in enumerate(headers):
            cell_value = (
                row.cells[col_idx].text.strip() if col_idx < len(row.cells) else ""
            )
            row_dict[header] = cell_value

        result.append(row_dict)

    return result
