"""Generate DOCX test fixture for docx_parser tests."""

from pathlib import Path

from docx import Document


def generate_process_docx():
    """Generate sample_process.docx with realistic PDD content."""
    output_path = (
        Path(__file__).parent.parent.parent
        / "data"
        / "sample_pdds"
        / "sample_process.docx"
    )

    doc = Document()

    # Heading 1: Process Overview
    doc.add_heading("Process Overview", level=1)

    doc.add_paragraph(
        "This process automates the extraction of daily sales data from SAP ECC "
        "and populates an Excel reporting template. The automation is built using "
        "Blue Prism and handles data validation and error scenarios."
    )

    # Heading 1: Target Applications
    doc.add_heading("Target Applications", level=1)
    doc.add_paragraph("SAP ECC", style="List Bullet")
    doc.add_paragraph("Excel", style="List Bullet")
    doc.add_paragraph("Outlook", style="List Bullet")

    # Heading 1: Process Steps
    doc.add_heading("Process Steps", level=1)

    doc.add_heading("Step 1 - Login to SAP", level=2)
    doc.add_paragraph(
        "The bot launches SAP GUI and navigates to the login screen. "
        "It enters credentials from a secure credential store and waits for "
        "the main menu to load before proceeding."
    )

    doc.add_heading("Step 2 - Navigate to Transaction", level=2)
    doc.add_paragraph(
        "Once logged in, the bot navigates to transaction VA05 to access the "
        "sales orders list. It applies filters for today's date and order status "
        "to retrieve the relevant transactions."
    )

    doc.add_heading("Step 3 - Extract Data", level=2)
    doc.add_paragraph(
        "The bot extracts detailed information from each sales order including "
        "order number, customer name, amount, and delivery date. Data is collected "
        "into an internal collection for processing."
    )

    doc.add_heading("Step 4 - Save to Excel", level=2)
    doc.add_paragraph(
        "The extracted data is written to an Excel spreadsheet using the "
        "'Daily_Sales_Report' and 'Sales_Summary' templates. The file is saved "
        "to the shared network drive with a timestamp."
    )

    # Heading 1: Business Rules
    doc.add_heading("Business Rules", level=1)
    doc.add_paragraph("Orders must be processed within the same business day")
    doc.add_paragraph("Amounts exceeding 50000 USD require additional validation")
    doc.add_paragraph("Failed records are logged and escalated to the operations team")

    # Heading 1: Digital Layouts
    doc.add_heading("Digital Layouts", level=1)
    doc.add_paragraph(
        "This process uses two key Excel templates: 'Daily_Sales_Report' for "
        "detailed transaction records and 'Sales_Summary' for aggregate reporting. "
        "Both templates include validation rules and conditional formatting."
    )

    # Add a table
    doc.add_heading("Process Mapping", level=1)
    table = doc.add_table(rows=5, cols=3)
    table.style = "Light Grid Accent 1"

    # Header row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Step"
    hdr_cells[1].text = "Application"
    hdr_cells[2].text = "Action"

    # Data rows
    data = [
        ("1", "SAP ECC", "Extract sales orders"),
        ("2", "Excel", "Validate data format"),
        ("3", "Excel", "Populate summary sheet"),
        ("4", "Outlook", "Send completion notification"),
    ]

    for row_idx, (step, app, action) in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        row_cells[0].text = step
        row_cells[1].text = app
        row_cells[2].text = action

    # Heading 1: Exception Handling
    doc.add_heading("Exception Handling", level=1)
    doc.add_paragraph(
        "If SAP connection fails, the bot retries up to three times with "
        "exponential backoff. If all retries fail, an alert email is sent to "
        "the RPA operations team. For data validation failures, records are "
        "logged to a quarantine file for manual review."
    )

    doc.save(str(output_path))
    print(f"Created {output_path}")


if __name__ == "__main__":
    output_dir = Path(__file__).parent.parent.parent / "data" / "sample_pdds"
    output_dir.mkdir(parents=True, exist_ok=True)
    generate_process_docx()
    print("DOCX fixture generated successfully!")
